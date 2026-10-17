from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django.db import transaction
from django.db.models import F
from rest_framework.permissions import IsAuthenticated
from code_editor.models import CodeDocument
from .models import (
    Document,
    Presentation,
    Slide,
    SlideElement,
    PresentationChat,
    ChatSummary
)

from pptx import Presentation as PPTXPresentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from django.http import HttpResponse

import time

from .serializers import (
    PresentationSerializer,
    SlideElementSerializer,
    SlideSerializer,
    PresentationFullSerializer,
    ChatSerializer
)

import google.generativeai as genai
from django.conf import settings

genai.configure(api_key=settings.GEMINI_API_KEY)


class PresentationViewSet(viewsets.ModelViewSet):
    serializer_class = PresentationSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        workspace_id = self.kwargs.get("workspace_pk")

        return Presentation.objects.filter(
            document__workspace_id=workspace_id,
            document__workspace__workspacemember__user=self.request.user
        )

    def perform_create(self, serializer):

        workspace_id = self.kwargs.get("workspace_pk")

        document_type = self.request.data.get("document_type", "ppt")

        document = Document.objects.create(
            workspace_id=workspace_id,
            title=self.request.data.get("title"),
            document_type=document_type,
            created_by=self.request.user
        )
        if document_type =="ppt":
            presentation = serializer.save(
                created_by=self.request.user,
                document=document
            )

            Slide.objects.create(
                presentation=presentation,
                order=0
            )
        elif document_type == "code":
            CodeDocument.objects.create(
                document=document,
                code=""
            )

    @action(detail=True, methods=["get"])
    def full(self, request, workspace_pk=None, pk=None):
        presentation = self.get_object()
        serializer = PresentationFullSerializer(presentation)
        return Response(serializer.data)

    @action(detail=True, methods=["get"])
    def export(self, request, workspace_pk=None, pk=None):
        start = time.time()

        presentation = self.get_object()

        prs = PPTXPresentation()

        slides = presentation.slides.all().order_by("order")

        for slide_obj in slides:

            slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(slide_layout)

            elements = slide_obj.elements.all().order_by("z_index")

            for element in elements:

                left = Inches(element.x / 96)
                top = Inches(element.y / 96)
                width = Inches(element.width / 96)
                height = Inches(element.height / 96)

                if element.type == "text":

                    textbox = slide.shapes.add_textbox(
                        left,
                        top,
                        width,
                        height
                    )

                    tf = textbox.text_frame
                    tf.text = element.content.get("text", "")

                    if tf.paragraphs and tf.paragraphs[0].runs:

                        run = tf.paragraphs[0].runs[0]
                        font = run.font

                        font_size = element.style.get("fontsize")

                        if font_size:
                            font.size = Pt(font_size)

                        color = element.style.get("color")

                        if color:
                            rgb = RGBColor.from_string(
                                color.replace("#", "")
                            )
                            font.color.rgb = rgb

                elif element.type == "rectangle":

                    shape = slide.shapes.add_shape(
                        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                        left,
                        top,
                        width,
                        height
                    )

                    fill_color = element.style.get("fillColor")

                    if fill_color:
                        rgb = RGBColor.from_string(
                            fill_color.replace("#", "")
                        )

                        shape.fill.solid()
                        shape.fill.fore_color.rgb = rgb

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )

        response["Content-Disposition"] = (
            f'attachment;filename="{presentation.title}.pptx"'
        )

        prs.save(response)

        end = time.time()

        print("Export took:", round(end - start, 3), "seconds")

        return response

    @action(detail=True, methods=["get"])
    def chat(self, request, workspace_pk=None, pk=None):

        chats = PresentationChat.objects.filter(
            presentation_id=pk
        ).order_by("created_at")

        serializer = ChatSerializer(chats, many=True)

        return Response(serializer.data)


class SlideViewSet(viewsets.ModelViewSet):

    serializer_class = SlideSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):

        presentation_id = self.kwargs.get("presentation_pk")

        return Slide.objects.filter(
            presentation_id=presentation_id,
            presentation__document__workspace__workspacemember__user=self.request.user
        )

    def perform_create(self, serializer):

        presentation = serializer.validated_data["presentation"]

        last_order = Slide.objects.filter(
            presentation=presentation
        ).count()

        serializer.save(order=last_order)

    @action(detail=True, methods=["patch"])
    def reorder(
        self,
        request,
        workspace_pk=None,
        presentation_pk=None,
        pk=None
    ):

        slide = self.get_object()

        new_order = request.data.get("new_order")

        if new_order is None:

            return Response(
                {"error": "new_order is Required"},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            new_order = int(new_order)

        except ValueError:

            return Response(
                {"error": "new_order must be integer"},
                status=status.HTTP_400_BAD_REQUEST
            )

        slides = Slide.objects.filter(
            presentation=slide.presentation
        ).order_by("order")

        max_order = slides.count() - 1

        if new_order < 0 or new_order > max_order:

            return Response(
                {"error": "Invalid order position"},
                status=status.HTTP_400_BAD_REQUEST
            )

        old_order = slide.order

        if new_order == old_order:
            return Response({"message": "No changes Needed"})

        with transaction.atomic():

            if new_order > old_order:

                Slide.objects.filter(
                    presentation=slide.presentation,
                    order__gt=old_order,
                    order__lte=new_order
                ).update(order=F("order") - 1)

            else:

                Slide.objects.filter(
                    presentation=slide.presentation,
                    order__gte=new_order,
                    order__lt=old_order
                ).update(order=F("order") + 1)

            slide.order = new_order
            slide.save()

        return Response({"message": "Slide reordered successfully"})


class SlideElementViewSet(viewsets.ModelViewSet):

    serializer_class = SlideElementSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):

        slide_id = self.kwargs.get("slide_pk")

        return SlideElement.objects.filter(
            slide_id=slide_id,
            slide__presentation__document__workspace__workspacemember__user=self.request.user
        )

    def perform_destroy(self, instance):

        z = instance.z_index
        slide = instance.slide

        instance.delete()

        SlideElement.objects.filter(
            slide=slide,
            z_index__gt=z
        ).update(z_index=F("z_index") - 1)


def generate_summary(presentation_id):

    chats = PresentationChat.objects.filter(
        presentation_id=presentation_id
    ).order_by("created_at")

    full_text = "\n".join([c.message for c in chats])

    try:
        summary = generate_ai_summary(full_text)

    except:
        summary = full_text[:500]

    ChatSummary.objects.update_or_create(
        presentation_id=presentation_id,
        defaults={"summary": summary}
    )


@action(detail=True, methods=["post"])
def generate_summary(self, request, pk=None):

    generate_summary(pk)

    return Response({
        "message": "Summary generated"
    })


def generate_ai_summary(text):

    model = genai.GenerativeModel("gemini-pro")

    prompt = f"""
    Summarize the following conversation into:
    - key points
    - decisions
    - action items

    Conversation:
    {text}
    """

    response = model.generate_content(prompt)

    return response.text